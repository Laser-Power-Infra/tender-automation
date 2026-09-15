import asyncio
import logging
import re
from pathlib import Path

from playwright.sync_api import sync_playwright
from django.conf import settings

from .browser import detect_chrome_path

if hasattr(asyncio, "WindowsProactorEventLoopPolicy"):
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

logger = logging.getLogger(__name__)


def _tender247_login(page, email: str, password: str) -> bool:
    print("[Tender247Result] Navigating to https://www.tender247.com/auth/tender ...")
    page.goto("https://www.tender247.com/auth/tender", timeout=50000)
    try:
        page.locator("input[name='emailId']").first.wait_for(state="visible", timeout=3000)
    except Exception:
        login_btn = page.get_by_role("button", name="Log in")
        if login_btn.count() == 0:
            login_btn = page.locator("button:has-text('Log in')")
        if login_btn.count() > 0:
            try:
                if login_btn.first.is_visible():
                    print("[Tender247Result] Clicking Log in...")
                    login_btn.first.click()
                    page.wait_for_timeout(1500)
            except Exception:
                pass
    signup_btn = page.locator("button:has-text('Sign Up')")
    if signup_btn.count() > 0:
        try:
            if signup_btn.first.is_visible():
                signup_btn.first.click()
                page.wait_for_timeout(1000)
        except Exception:
            pass
    page.locator("input[name='emailId']").first.wait_for(state="visible", timeout=10000)
    page.locator("input[name='emailId']").first.fill(email)
    page.locator("input[name='password']").first.fill(password)
    page.locator("button[type='submit']:has-text('Submit')").click()
    page.wait_for_timeout(5000)
    close_btn = page.locator("button:has(span.sr-only:text('Close'))")
    if close_btn.count() > 0:
        try:
            if close_btn.first.is_visible():
                print("[Tender247Result] Closing dialog...")
                close_btn.first.click()
                page.wait_for_timeout(1000)
        except Exception:
            pass
    # ponytail: same success heuristic as non_gem_tender_pdf_downloader login_tender247
    try:
        signup_text = page.locator("button:has-text('Sign Up')").inner_text() if page.locator("button:has-text('Sign Up')").count() > 0 else ""
    except Exception:
        signup_text = ""
    return "Sign Up" not in signup_text


def _click_indian_result(page) -> None:
    print("[Tender247Result] Clicking Result menu...")
    result_loc = page.locator("li.cursor-pointer:has(img[src*='tender-result-icon']):has-text('Result')")
    result_loc.first.wait_for(state="visible", timeout=15000)
    result_loc.first.scroll_into_view_if_needed(timeout=5000)
    result_loc.first.click()
    page.wait_for_timeout(1500)
    try:
        page.wait_for_load_state("networkidle", timeout=10000)
    except Exception:
        pass
    print("[Tender247Result] Clicking Indian...")
    # ponytail: exact href scopes away from global Indian text dupes
    candidates = [
        "a[href='/auth/result']:has-text('Indian')",
        "a[href='/auth/result']",
        "a:has(img[src*='domestic-icon']):has-text('Indian')",
        "a:has-text('Indian')",
    ]
    clicked = False
    for sel in candidates:
        loc = page.locator(sel).first
        try:
            if loc.count() > 0 and loc.is_visible():
                print(f"[Tender247Result] Indian via {sel!r}")
                loc.click()
                clicked = True
                break
        except Exception:
            continue
    if not clicked:
        loc = page.get_by_text("Indian", exact=True).first
        loc.wait_for(state="visible", timeout=15000)
        loc.click()
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass


def _click_today_results(page) -> None:
    print("[Tender247Result] Clicking Today Results...")
    # ponytail: card has img alt Today Results + p text — scopes away from Today Tenders
    loc = page.locator("div.cursor-pointer:has(p:has-text('Today Results'))")
    if loc.count() == 0:
        loc = page.locator("div.cursor-pointer:has(img[alt='Today Results'])")
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    loc.first.click()
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass


def _download_today_excel(page) -> Path:
    print("[Tender247Result] Waiting Today Results data...")
    page.wait_for_timeout(2000)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass
    print("[Tender247Result] Clicking Download Excel...")
    loc = page.locator("span:has(img[alt='Download Excel'])")
    if loc.count() == 0:
        loc = page.locator("img[alt='Download Excel']")
    loc.first.wait_for(state="visible", timeout=15000)
    loc.first.scroll_into_view_if_needed(timeout=5000)
    download_dir = Path(settings.TENDER_PARSING_TEMP_DIR)
    download_dir.mkdir(parents=True, exist_ok=True)
    with page.expect_download(timeout=30000) as dl:
        try:
            loc.first.click()
        except Exception:
            loc.first.evaluate("el => el.click()")
    download = dl.value
    suggested = download.suggested_filename or "tender247_today.xlsx"
    dest = download_dir / suggested
    download.save_as(str(dest))
    print(f"[Tender247Result] Downloaded: {dest}")
    if not dest.exists() or dest.stat().st_size == 0:
        raise RuntimeError(f"Tender247 excel not captured: {dest}")
    return dest


def _parse_247_excel(path: Path) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return []
    headers: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(c).strip() if c is not None else "" for c in row]
        break
    headers = [h for h in headers if h]
    wb.close()
    logger.info("[Tender247Result] Excel headers: %s", headers)
    print(f"[Tender247Result] Excel headers: {headers}")
    return headers


def _iter_247_rows(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return
    try:
        for row in ws.iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def _map_stage(raw: str) -> str:
    # ponytail: substring match on Tender Stage → normalized status
    low = (raw or "").strip().lower()
    if re.search(r"\baoc\b", low):
        return "AWARDED"
    if "financial" in low:
        return "FINANCIAL EVALUATION"
    if "technical" in low:
        return "TECHNICAL BID OPENED"
    return ""


def _log_247_updates(path: Path) -> list[dict]:
    from .tender_result_updater import update_tender_result

    gen = _iter_247_rows(path)
    try:
        headers_raw = next(gen)
    except StopIteration:
        return []
    headers = [str(c).strip() if c is not None else "" for c in headers_raw]
    hmap = {h.lower(): idx for idx, h in enumerate(headers)}
    idx_ref = hmap.get("tender reference no")
    idx_l1 = hmap.get("winner bidder")
    idx_amt = hmap.get("contract value")
    idx_comp = hmap.get("participator bidders")
    idx_status = hmap.get("tender stage")
    if idx_ref is None or idx_l1 is None or idx_amt is None:
        logger.warning("[Tender247Result] missing columns headers=%s", headers)
        print(f"[Tender247Result] missing columns headers={headers}")
        return []
    updates: list[dict] = []
    for r in gen:
        if not r:
            continue
        ref = str(r[idx_ref]).strip() if r[idx_ref] is not None else ""
        if not ref:
            continue
        l1 = str(r[idx_l1]).strip() if r[idx_l1] is not None else ""
        amt_raw = str(r[idx_amt]).strip() if r[idx_amt] is not None else ""
        competitors = str(r[idx_comp]).strip() if idx_comp is not None and r[idx_comp] is not None else ""
        cur_raw = str(r[idx_status]).strip() if idx_status is not None and r[idx_status] is not None else ""
        cur_status = _map_stage(cur_raw)
        res = update_tender_result(ref, l1, amt_raw, competitors or None, cur_status or None)
        updates.append(res)
    logger.info("[Tender247Result] total rows %d", len(updates))
    print(f"[Tender247Result] total rows {len(updates)}")
    return updates


def get_tender247_result(email: str | None = None, password: str | None = None) -> dict:
    email = email or settings.TENDER247_EMAIL
    password = password or settings.TENDER247_PASSWORD
    if not email or not password:
        return {"success": False, "error": "TENDER247_EMAIL/PASSWORD not configured"}

    chrome_path = detect_chrome_path()
    tenders_file: Path | None = None
    url = ""
    # Phase A: Playwright
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=chrome_path, headless=settings.HEADLESS_BROWSER)
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        try:
            ok = _tender247_login(page, email, password)
            if not ok:
                return {"success": False, "error": "Tender247 login failed", "url": page.url}
            _click_indian_result(page)
            _click_today_results(page)
            tenders_file = _download_today_excel(page)
            url = page.url
        except Exception as e:
            logger.exception("[Tender247Result] failed (download): %s", e)
            return {"success": False, "error": str(e), "url": page.url if 'page' in locals() else ""}
        finally:
            try:
                browser.close()
            except Exception:
                pass
    # clear async loop before DB like tiger
    try:
        asyncio.set_event_loop(None)
    except Exception:
        pass
    try:
        from django.db import close_old_connections
        close_old_connections()
    except Exception:
        pass
    # Phase B: DB
    try:
        assert tenders_file is not None
        headers = _parse_247_excel(tenders_file)
        updates = _log_247_updates(tenders_file)
        return {"success": True, "url": url, "tenders_file": str(tenders_file), "headers": headers, "updates": updates, "would_update": sum(1 for u in updates if u["found"])}
    except Exception as e:
        logger.exception("[Tender247Result] failed (DB): %s", e)
        return {"success": False, "error": str(e)}
